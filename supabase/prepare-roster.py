"""Read the roster and build a private SQL import; never changes the workbook.

Usage: python prepare-roster.py INPUT.xlsx PRIVATE_OUTPUT_DIRECTORY
Requires openpyxl for read-only workbook extraction.
"""
import csv
import json
import re
import sys
import uuid
from pathlib import Path

import openpyxl


def clean(value):
    return str(value or '').strip()


def key(value):
    return clean(value).casefold()


def quote(value):
    return "'" + str(value).replace("'", "''") + "'"


def prepare(workbook):
    team_rows = list(workbook['Teams'].iter_rows(values_only=True))
    scopes = {(clean(r[0]), clean(r[1])) for r in team_rows[1:] if r and r[0]}
    memberships, review = {}, []
    people = {}
    for row_number, raw in enumerate(workbook['Members'].iter_rows(values_only=True), 1):
        if row_number == 1 or not any(raw):
            continue
        row = list(raw) + [None] * max(0, 9 - len(raw))
        _, email, name, team, subteam, role, github, _, note = map(clean, row[:9])
        email = email.lower()
        item = dict(source_row=row_number, name=name, email=email, team=team,
                    subteam=subteam, note=note)
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email) or not name or not team:
            review.append(dict(item, reason='Missing confirmed person, email, or group'))
            continue
        roots = [s for s in scopes if key(s[0]) == key(team) and not s[1]]
        if len(roots) != 1:
            review.append(dict(item, reason='Unknown or ambiguous parent team'))
            continue
        targets = [roots[0]]
        if subteam:
            matches = [s for s in scopes if key(s[1]) == key(subteam)]
            local = [s for s in matches if key(s[0]) == key(team)]
            matches = local or matches
            if len(matches) != 1:
                review.append(dict(item, reason='Unknown or ambiguous subteam'))
                continue
            # A leaf membership must not grant every sibling group. Cross-team
            # membership retains the explicitly listed root as well as the leaf.
            targets = matches if local else targets + matches
            if not local:
                review.append(dict(item, reason='Resolved cross-team subteam from Teams tab; both memberships included'))
        # Group workbook aliases into one roster person. Login matching remains
        # exclusively by verified email, never by this name/GitHub key.
        person = key(github) or key(name)
        people.setdefault(person, {'name': name, 'emails': set()})['emails'].add(email)
        for parent, child in targets:
            memberships.setdefault((email, parent, child), dict(
                email=email, team=parent, subteam=child, source_row=row_number))
    return list(memberships.values()), review, [dict(name=p['name'], key=k, emails=sorted(p['emails']))
                                              for k, p in people.items()]


def people_sql(people):
    """Link explicit workbook aliases to stable people; never match auth by name."""
    statements = ['do $people$\ndeclare existing_ids uuid[]; target uuid;\nbegin']
    for person in people:
        person_id = str(uuid.uuid5(uuid.NAMESPACE_URL, 'exo-workbench-roster:' + person['key']))
        emails = ','.join(quote(email) for email in person['emails'])
        statements.append(f"""
  select array_agg(distinct person_id) into existing_ids from exo_private.roster_memberships
    where email in ({emails}) and person_id is not null;
  if coalesce(cardinality(existing_ids),0)>1 then
    raise exception 'Conflicting existing person mappings. Review aliases before merging.';
  end if;
  target := coalesce(existing_ids[1], {quote(person_id)}::uuid);
  insert into public.roster_people(id,display_name) values(target,{quote(person['name'])})
    on conflict(id) do nothing;
  update exo_private.roster_memberships set person_id=target where email in ({emails})
    and person_id is null;
""")
    statements.append('end;\n$people$;\n')
    return '\n'.join(statements)


def seed_sql(records):
    if not records:
        raise ValueError('No unambiguous memberships found')
    values = ',\n'.join('  (' + ', '.join([quote(r['email']), quote(r['team']),
                       quote(r['subteam']), str(r['source_row'])]) + ')' for r in records)
    return """
-- PRIVATE roster import. Run in Supabase SQL Editor as postgres.
-- Both listed aliases are approved; matching is by verified Google email only.
-- No leadership roles are assigned. Existing memberships are never removed.
begin;
set local standard_conforming_strings = on;
-- A single statement owns the entire seed. No temporary tables or session state.
do $import$
declare
  entry record;
  parent_ids uuid[];
  child_ids uuid[];
  parent_id uuid;
  target_id uuid;
  account_id uuid;
begin
  lock table public.teams in share row exclusive mode;
  for entry in select * from (values
""" + values + """
  ) as input(email, parent_name, child_name, source_row)
  loop
    select array_agg(t.id) into parent_ids from public.teams t
      where t.parent_team_id is null
        and lower(btrim(t.name)) = lower(btrim(entry.parent_name));
    if coalesce(cardinality(parent_ids), 0) > 1 then
      raise exception 'Duplicate parent team: %. Import rolled back.', entry.parent_name;
    elsif coalesce(cardinality(parent_ids), 0) = 0 then
      insert into public.teams(name) values(entry.parent_name) returning id into parent_id;
    else parent_id := parent_ids[1]; end if;
    target_id := parent_id;
    if entry.child_name <> '' then
      select array_agg(t.id) into child_ids from public.teams t
        where t.parent_team_id = parent_id
          and lower(btrim(t.name)) = lower(btrim(entry.child_name));
      if coalesce(cardinality(child_ids), 0) > 1 then
        raise exception 'Duplicate subteam: % / %. Import rolled back.', entry.parent_name, entry.child_name;
      elsif coalesce(cardinality(child_ids), 0) = 0 then
        insert into public.teams(name, parent_team_id)
          values(entry.child_name, parent_id) returning id into target_id;
      else target_id := child_ids[1]; end if;
    end if;
    insert into exo_private.roster_memberships(email, team_id, source, source_row)
      values(entry.email, target_id, 'Workbook roster import', entry.source_row)
      on conflict(email, team_id) do nothing;

    -- Existing verified accounts receive their pending approvals immediately.
    -- Previously consumed/disabled approvals remain untouched on reruns.
    for account_id in select u.id from auth.users u
      where lower(btrim(u.email)) = entry.email order by u.id
    loop perform exo_private.apply_roster_memberships(account_id); end loop;
  end loop;
end;
$import$;

-- Standalone verification also works after reconnecting to SQL Editor.
with input(email, parent_name, child_name, source_row) as (values
""" + values + """
)
select r.email, r.parent_name as team, nullif(r.child_name,'') as subteam,
  a.claimed_by as user_id,
  case when a.disabled_at is not null then 'Disabled approval'
       when a.claimed_at is null then 'Waiting for verified Google sign-in'
       when exists(select 1 from public.team_members m
         where m.team_id=a.team_id and m.user_id=a.claimed_by) then 'Assigned'
       else 'Previously consumed; membership subsequently removed' end as enrollment_status
from input r
join public.teams p on p.parent_team_id is null
  and lower(btrim(p.name))=lower(btrim(r.parent_name))
join public.teams t on (r.child_name='' and t.id=p.id)
  or (r.child_name<>'' and t.parent_team_id=p.id
    and lower(btrim(t.name))=lower(btrim(r.child_name)))
join exo_private.roster_memberships a on a.email=r.email and a.team_id=t.id
order by r.parent_name, r.child_name, r.email;
commit;
"""


def main():
    source, output = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
    repo = Path(__file__).resolve().parent.parent
    if output == repo or repo in output.parents:
        raise ValueError('Private roster output must be outside the website/repository')
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    records, review, people = prepare(workbook)
    workbook.close()
    migration = Path(__file__).with_name('roster-memberships.sql').read_text(encoding='utf-8')
    # Install triggers and seed in one transaction so signups cannot fall between.
    migration = migration.replace('begin;', '', 1).replace('commit;', '', 1)
    assignment_migration = Path(__file__).with_name('roster-task-assignment.sql').read_text(encoding='utf-8')
    assignment_migration = assignment_migration.replace('begin;', '', 1).rsplit('commit;', 1)[0]
    seed = seed_sql(records)
    combined = seed.replace('begin;', 'begin;\n' + migration + '\n' + assignment_migration, 1)
    combined = combined.rsplit('commit;', 1)[0] + people_sql(people) + 'commit;\n'
    output.mkdir(parents=True, exist_ok=True)
    (output / 'install-roster.sql').write_text(combined, encoding='utf-8')
    report = dict(source=str(source), people=len(people), email_accounts=len({r['email'] for r in records}),
                  memberships=len(records), aliases=[p for p in people if len(p['emails']) > 1],
                  review=review)
    (output / 'roster-review.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    with (output / 'approved-memberships.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['email', 'team', 'subteam', 'source_row'])
        writer.writeheader()
        writer.writerows(records)
    print(json.dumps({k: report[k] for k in ('people', 'email_accounts', 'memberships')}))
    print('Review entries:', len(review))
    print('Private output:', output)


if __name__ == '__main__':
    main()
